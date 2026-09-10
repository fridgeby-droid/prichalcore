from __future__ import annotations

from datetime import datetime, timedelta
from typing import Any
from urllib.parse import quote
from io import BytesIO, StringIO
import csv

import httpx
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from pydantic import BaseModel, Field
from sqlalchemy import delete, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from starlette.responses import StreamingResponse

from app.core.config import PHOTO_REQUEST_TTL_MINUTES
from app.core.security import get_current_user
from app.core.serializers import row
from app.db.database import get_db
from app.db.models import (
    Category,
    OrderScheduleRule,
    PriceList,
    PriceListUploadRequest,
    Product,
    ProductOrderLimit,
    Store,
    StoreProduct,
    Supplier,
    UnitOfMeasure,
)
from app.services.telegram import file_download_url, get_file_path, send_message

router = APIRouter(prefix="/api/order-settings", tags=["order-settings"])


def _commit_or_conflict(db: Session, message: str = "Объект используется и не может быть удалён"):
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, message)


# Permissions are intentionally open to every authenticated active user in v1.1.
# The Roles editor will replace this temporary policy later.


class StoreIn(BaseModel):
    name: str
    code: str | None = None


@router.get("/stores")
def stores(user=Depends(get_current_user), db: Session = Depends(get_db)):
    return [row(x, "id", "name", "code", "active") for x in db.scalars(select(Store).order_by(Store.name)).all()]


@router.post("/stores")
def create_store(payload: StoreIn, user=Depends(get_current_user), db: Session = Depends(get_db)):
    obj = Store(name=payload.name.strip(), code=(payload.code or None))
    db.add(obj); _commit_or_conflict(db, "Магазин с таким названием/кодом уже существует"); db.refresh(obj)
    return row(obj, "id", "name", "code", "active")


@router.patch("/stores/{store_id}")
def update_store(store_id: int, payload: dict[str, Any], user=Depends(get_current_user), db: Session = Depends(get_db)):
    obj = db.get(Store, store_id)
    if not obj: raise HTTPException(404, "Магазин не найден")
    for k in ("name", "code", "active"):
        if k in payload: setattr(obj, k, payload[k])
    _commit_or_conflict(db, "Не удалось изменить магазин")
    return row(obj, "id", "name", "code", "active")


@router.delete("/stores/{store_id}")
def delete_store(store_id: int, user=Depends(get_current_user), db: Session = Depends(get_db)):
    obj = db.get(Store, store_id)
    if not obj: raise HTTPException(404, "Магазин не найден")
    db.delete(obj); _commit_or_conflict(db, "Магазин уже используется. Отключите его вместо удаления.")
    return {"ok": True}


class SupplierIn(BaseModel):
    name: str
    deadline_time: str = "10:00"


@router.get("/suppliers")
def suppliers(user=Depends(get_current_user), db: Session = Depends(get_db)):
    return [row(x, "id", "name", "deadline_time", "active") for x in db.scalars(select(Supplier).order_by(Supplier.name)).all()]


@router.post("/suppliers")
def create_supplier(payload: SupplierIn, user=Depends(get_current_user), db: Session = Depends(get_db)):
    obj = Supplier(name=payload.name.strip(), deadline_time=payload.deadline_time)
    db.add(obj); _commit_or_conflict(db, "Поставщик с таким названием уже существует"); db.refresh(obj)
    return row(obj, "id", "name", "deadline_time", "active")


@router.patch("/suppliers/{supplier_id}")
def update_supplier(supplier_id: int, payload: dict[str, Any], user=Depends(get_current_user), db: Session = Depends(get_db)):
    obj = db.get(Supplier, supplier_id)
    if not obj: raise HTTPException(404, "Поставщик не найден")
    for k in ("name", "deadline_time", "active"):
        if k in payload: setattr(obj, k, payload[k])
    _commit_or_conflict(db, "Не удалось изменить поставщика")
    return row(obj, "id", "name", "deadline_time", "active")


@router.delete("/suppliers/{supplier_id}")
def delete_supplier(supplier_id: int, user=Depends(get_current_user), db: Session = Depends(get_db)):
    obj = db.get(Supplier, supplier_id)
    if not obj: raise HTTPException(404, "Поставщик не найден")
    db.delete(obj); _commit_or_conflict(db, "Поставщик уже используется. Отключите его вместо удаления.")
    return {"ok": True}


class CategoryIn(BaseModel):
    name: str
    supplier_id: int | None = None


@router.get("/categories")
def categories(user=Depends(get_current_user), db: Session = Depends(get_db)):
    return [row(x, "id", "name", "supplier_id", "active") for x in db.scalars(select(Category).order_by(Category.name)).all()]


@router.post("/categories")
def create_category(payload: CategoryIn, user=Depends(get_current_user), db: Session = Depends(get_db)):
    obj = Category(name=payload.name.strip(), supplier_id=payload.supplier_id)
    db.add(obj); _commit_or_conflict(db, "Такая категория уже существует"); db.refresh(obj)
    return row(obj, "id", "name", "supplier_id", "active")


@router.patch("/categories/{category_id}")
def update_category(category_id: int, payload: dict[str, Any], user=Depends(get_current_user), db: Session = Depends(get_db)):
    obj = db.get(Category, category_id)
    if not obj: raise HTTPException(404, "Категория не найдена")
    for k in ("name", "supplier_id", "active"):
        if k in payload: setattr(obj, k, payload[k])
    _commit_or_conflict(db, "Не удалось изменить категорию")
    return row(obj, "id", "name", "supplier_id", "active")


@router.delete("/categories/{category_id}")
def delete_category(category_id: int, user=Depends(get_current_user), db: Session = Depends(get_db)):
    obj = db.get(Category, category_id)
    if not obj: raise HTTPException(404, "Категория не найдена")
    db.delete(obj); _commit_or_conflict(db, "Категория используется товарами. Отключите её вместо удаления.")
    return {"ok": True}


class UnitIn(BaseModel):
    name: str
    short_name: str


@router.get("/units")
def units(user=Depends(get_current_user), db: Session = Depends(get_db)):
    return [row(x, "id", "name", "short_name", "active") for x in db.scalars(select(UnitOfMeasure).order_by(UnitOfMeasure.name)).all()]


@router.post("/units")
def create_unit(payload: UnitIn, user=Depends(get_current_user), db: Session = Depends(get_db)):
    obj = UnitOfMeasure(name=payload.name.strip(), short_name=payload.short_name.strip())
    db.add(obj); _commit_or_conflict(db, "Такая единица измерения уже существует"); db.refresh(obj)
    return row(obj, "id", "name", "short_name", "active")


@router.patch("/units/{unit_id}")
def update_unit(unit_id: int, payload: dict[str, Any], user=Depends(get_current_user), db: Session = Depends(get_db)):
    obj = db.get(UnitOfMeasure, unit_id)
    if not obj: raise HTTPException(404, "Единица измерения не найдена")
    for k in ("name", "short_name", "active"):
        if k in payload: setattr(obj, k, payload[k])
    _commit_or_conflict(db, "Не удалось изменить единицу измерения")
    return row(obj, "id", "name", "short_name", "active")


@router.delete("/units/{unit_id}")
def delete_unit(unit_id: int, user=Depends(get_current_user), db: Session = Depends(get_db)):
    obj = db.get(UnitOfMeasure, unit_id)
    if not obj: raise HTTPException(404, "Единица измерения не найдена")
    db.delete(obj); _commit_or_conflict(db)
    return {"ok": True}


class ProductIn(BaseModel):
    name: str
    supplier_id: int
    category_id: int | None = None
    unit: str = "шт"


@router.get("/products")
def products(user=Depends(get_current_user), db: Session = Depends(get_db)):
    return [row(x, "id", "name", "supplier_id", "category_id", "unit", "active") for x in db.scalars(select(Product).order_by(Product.name)).all()]


@router.post("/products")
def create_product(payload: ProductIn, user=Depends(get_current_user), db: Session = Depends(get_db)):
    obj = Product(name=payload.name.strip(), supplier_id=payload.supplier_id, category_id=payload.category_id, unit=payload.unit.strip())
    db.add(obj); _commit_or_conflict(db, "Такой товар уже есть у поставщика"); db.refresh(obj)
    return row(obj, "id", "name", "supplier_id", "category_id", "unit", "active")


@router.patch("/products/{product_id}")
def update_product(product_id: int, payload: dict[str, Any], user=Depends(get_current_user), db: Session = Depends(get_db)):
    obj = db.get(Product, product_id)
    if not obj: raise HTTPException(404, "Товар не найден")
    for k in ("name", "supplier_id", "category_id", "unit", "active"):
        if k in payload: setattr(obj, k, payload[k])
    _commit_or_conflict(db, "Не удалось изменить товар")
    return row(obj, "id", "name", "supplier_id", "category_id", "unit", "active")


@router.delete("/products/{product_id}")
def delete_product(product_id: int, user=Depends(get_current_user), db: Session = Depends(get_db)):
    obj = db.get(Product, product_id)
    if not obj: raise HTTPException(404, "Товар не найден")
    db.delete(obj); _commit_or_conflict(db, "Товар уже используется в заявках. Отключите его вместо удаления.")
    return {"ok": True}


class AssortmentIn(BaseModel):
    store_id:int
    product_id:int
    enabled:bool=True


@router.get("/assortment")
def assortment(store_id:int,user=Depends(get_current_user),db:Session=Depends(get_db)):
    mappings=list(db.scalars(select(StoreProduct).where(StoreProduct.store_id==store_id)).all())
    enabled={x.product_id for x in mappings if x.enabled}
    explicit=bool(mappings)
    products=list(db.scalars(select(Product).where(Product.active.is_(True)).order_by(Product.name)).all())
    return [{**row(p,"id","name","supplier_id","category_id","unit"),"enabled":(p.id in enabled if explicit else True)} for p in products]


@router.put("/assortment")
def save_assortment(payload:AssortmentIn,user=Depends(get_current_user),db:Session=Depends(get_db)):
    existing=list(db.scalars(select(StoreProduct).where(StoreProduct.store_id==payload.store_id)).all())
    if not existing:
        for pid in db.scalars(select(Product.id).where(Product.active.is_(True))).all():
            db.add(StoreProduct(store_id=payload.store_id,product_id=pid,enabled=True))
        db.flush()
    obj=db.get(StoreProduct,{"store_id":payload.store_id,"product_id":payload.product_id})
    if not obj:
        obj=StoreProduct(store_id=payload.store_id,product_id=payload.product_id,enabled=payload.enabled);db.add(obj)
    else:obj.enabled=payload.enabled
    db.commit();return {"ok":True}


class LimitIn(BaseModel):
    product_id: int
    store_id: int | None = None
    min_qty: float | None = None
    max_qty: float | None = None
    step_qty: float | None = None
    active: bool = True


@router.get("/limits")
def limits(user=Depends(get_current_user), db: Session = Depends(get_db)):
    result=[]
    q=select(ProductOrderLimit, Product, Store).join(Product, Product.id==ProductOrderLimit.product_id).outerjoin(Store, Store.id==ProductOrderLimit.store_id).order_by(Product.name)
    for lim,p,s in db.execute(q).all():
        result.append({
            **row(lim,"id","product_id","store_id","active"),
            "product_name":p.name,"store_name":s.name if s else "Все магазины",
            "min_qty":float(lim.min_qty) if lim.min_qty is not None else None,
            "max_qty":float(lim.max_qty) if lim.max_qty is not None else None,
            "step_qty":float(lim.step_qty) if lim.step_qty is not None else None,
        })
    return result


@router.post("/limits")
def create_limit(payload: LimitIn, user=Depends(get_current_user), db: Session = Depends(get_db)):
    q=select(ProductOrderLimit).where(ProductOrderLimit.product_id==payload.product_id, ProductOrderLimit.store_id==payload.store_id)
    obj=db.scalar(q)
    if obj:
        for k,v in payload.model_dump().items(): setattr(obj,k,v)
    else:
        obj=ProductOrderLimit(**payload.model_dump());db.add(obj)
    _commit_or_conflict(db, "Не удалось сохранить лимит"); db.refresh(obj)
    return {"id":obj.id,"ok":True}


@router.patch("/limits/{limit_id}")
def update_limit(limit_id:int,payload:dict[str,Any],user=Depends(get_current_user),db:Session=Depends(get_db)):
    obj=db.get(ProductOrderLimit,limit_id)
    if not obj:raise HTTPException(404,"Лимит не найден")
    for k in ("product_id","store_id","min_qty","max_qty","step_qty","active"):
        if k in payload:setattr(obj,k,payload[k])
    _commit_or_conflict(db,"Не удалось изменить лимит")
    return {"ok":True}


@router.delete("/limits/{limit_id}")
def delete_limit(limit_id:int,user=Depends(get_current_user),db:Session=Depends(get_db)):
    obj=db.get(ProductOrderLimit,limit_id)
    if not obj:raise HTTPException(404,"Лимит не найден")
    db.delete(obj);db.commit();return {"ok":True}


class ScheduleIn(BaseModel):
    store_id:int
    supplier_id:int
    weekdays:list[int]


@router.get("/schedule")
def schedule(user=Depends(get_current_user),db:Session=Depends(get_db)):
    stores=[row(x,"id","name","active") for x in db.scalars(select(Store).order_by(Store.name)).all()]
    suppliers=[row(x,"id","name","active") for x in db.scalars(select(Supplier).order_by(Supplier.name)).all()]
    rules=[row(x,"id","store_id","supplier_id","schedule_type","weekday","active") for x in db.scalars(select(OrderScheduleRule).where(OrderScheduleRule.active.is_(True),OrderScheduleRule.schedule_type=="order").order_by(OrderScheduleRule.store_id,OrderScheduleRule.weekday)).all()]
    return {"stores":stores,"suppliers":suppliers,"rules":rules}


@router.put("/schedule")
def save_schedule(payload:ScheduleIn,user=Depends(get_current_user),db:Session=Depends(get_db)):
    days=sorted({int(x) for x in payload.weekdays if 0<=int(x)<=6})
    db.execute(delete(OrderScheduleRule).where(OrderScheduleRule.store_id==payload.store_id,OrderScheduleRule.supplier_id==payload.supplier_id,OrderScheduleRule.schedule_type=="order"))
    for day in days:
        db.add(OrderScheduleRule(store_id=payload.store_id,supplier_id=payload.supplier_id,schedule_type="order",weekday=day,active=True))
    db.commit();return {"ok":True,"weekdays":days}



def _norm_header(value: Any) -> str:
    return str(value or "").strip().lower().replace("ё", "е")


def _first_value(data: dict[str, Any], aliases: tuple[str, ...]):
    for alias in aliases:
        if alias in data and str(data[alias] or "").strip():
            return data[alias]
    return None


def _parse_product_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    lower = (filename or "").lower()
    rows: list[dict[str, Any]] = []
    if lower.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(500, "На сервере не установлен openpyxl") from exc
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        values = list(ws.iter_rows(values_only=True))
        if not values:
            return []
        headers = [_norm_header(x) for x in values[0]]
        for vals in values[1:]:
            data = {headers[i]: vals[i] if i < len(vals) else None for i in range(len(headers)) if headers[i]}
            if any(v not in (None, "") for v in data.values()):
                rows.append(data)
        return rows
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = content.decode("cp1251")
        except UnicodeDecodeError as exc:
            raise HTTPException(400, "CSV должен быть в UTF-8 или Windows-1251") from exc
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if ";" in sample else ("\t" if "\t" in sample else ",")
    reader = csv.DictReader(StringIO(text), delimiter=delimiter)
    for raw in reader:
        data = {_norm_header(k): v for k, v in raw.items() if k is not None}
        if any(str(v or "").strip() for v in data.values()):
            rows.append(data)
    return rows


@router.get("/products/import-template")
def product_import_template(user=Depends(get_current_user)):
    content = "Товар;Поставщик;Категория;Единица\nПример товара;Пример поставщика;Пример категории;шт\n"
    return StreamingResponse(
        iter([content.encode("utf-8-sig")]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=products_import_template.csv"},
    )


@router.post("/products/import")
async def import_products(file: UploadFile = File(...), user=Depends(get_current_user), db: Session = Depends(get_db)):
    content = await file.read()
    if not content:
        raise HTTPException(400, "Файл пуст")
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(413, "Файл больше 10 МБ")
    filename = file.filename or "products.csv"
    if not filename.lower().endswith((".csv", ".tsv", ".txt", ".xlsx")):
        raise HTTPException(400, "Поддерживаются XLSX, CSV и TSV")
    parsed = _parse_product_rows(filename, content)
    if not parsed:
        raise HTTPException(400, "В таблице нет строк с товарами")

    suppliers = list(db.scalars(select(Supplier)).all())
    supplier_map = {x.name.strip().lower(): x for x in suppliers}
    categories = list(db.scalars(select(Category)).all())
    category_map = {(x.supplier_id, x.name.strip().lower()): x for x in categories}
    units = list(db.scalars(select(UnitOfMeasure)).all())
    unit_map = {x.short_name.strip().lower(): x for x in units}
    unit_map.update({x.name.strip().lower(): x for x in units})
    products = list(db.scalars(select(Product)).all())
    product_map = {(x.supplier_id, x.name.strip().lower()): x for x in products}

    inserted = 0
    updated = 0
    errors: list[dict[str, Any]] = []
    for idx, data in enumerate(parsed, start=2):
        name = str(_first_value(data, ("товар", "название", "наименование", "product", "name")) or "").strip()
        supplier_name = str(_first_value(data, ("поставщик", "supplier")) or "").strip()
        category_name = str(_first_value(data, ("категория", "category")) or "").strip()
        unit_name = str(_first_value(data, ("единица", "ед. изм.", "ед изм", "единица измерения", "unit")) or "шт").strip()
        if not name or not supplier_name:
            errors.append({"row": idx, "error": "Нужны столбцы Товар и Поставщик"})
            continue
        supplier = supplier_map.get(supplier_name.lower())
        if not supplier:
            errors.append({"row": idx, "error": f"Поставщик «{supplier_name}» не найден"})
            continue

        category_id = None
        if category_name:
            ckey = (supplier.id, category_name.lower())
            category = category_map.get(ckey)
            if not category:
                category = Category(name=category_name, supplier_id=supplier.id, active=True)
                db.add(category)
                db.flush()
                category_map[ckey] = category
            category_id = category.id

        unit = unit_map.get(unit_name.lower())
        if not unit:
            unit = UnitOfMeasure(name=unit_name, short_name=unit_name, active=True)
            db.add(unit)
            db.flush()
            unit_map[unit_name.lower()] = unit

        pkey = (supplier.id, name.lower())
        product = product_map.get(pkey)
        if product:
            product.category_id = category_id
            product.unit = unit.short_name
            product.active = True
            updated += 1
        else:
            product = Product(name=name, supplier_id=supplier.id, category_id=category_id, unit=unit.short_name, active=True)
            db.add(product)
            db.flush()
            product_map[pkey] = product
            inserted += 1
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, "Импорт не завершён из-за дублирующихся данных") from exc
    return {"ok": True, "rows": len(parsed), "inserted": inserted, "updated": updated, "errors": errors[:100]}


class PriceListRequestIn(BaseModel):
    supplier_id:int
    title:str|None=None


@router.get("/price-lists")
def price_lists(user=Depends(get_current_user),db:Session=Depends(get_db)):
    q=select(PriceList,Supplier).join(Supplier,Supplier.id==PriceList.supplier_id).where(PriceList.active.is_(True)).order_by(PriceList.created_at.desc())
    return [{**row(x,"id","supplier_id","title","file_name","mime_type","file_size","created_at"),"supplier_name":s.name} for x,s in db.execute(q).all()]


@router.post("/price-lists/request")
def request_price_list(payload:PriceListRequestIn,user=Depends(get_current_user),db:Session=Depends(get_db)):
    supplier=db.get(Supplier,payload.supplier_id)
    if not supplier:raise HTTPException(404,"Поставщик не найден")
    for old in db.scalars(select(PriceListUploadRequest).where(PriceListUploadRequest.user_id==user.id,PriceListUploadRequest.status=="waiting")).all():old.status="cancelled"
    req=PriceListUploadRequest(user_id=user.id,supplier_id=payload.supplier_id,title=payload.title or f"Прайс-лист {supplier.name}",status="waiting",expires_at=datetime.utcnow()+timedelta(minutes=PHOTO_REQUEST_TTL_MINUTES))
    db.add(req);db.commit();db.refresh(req)
    send_message(user.telegram_id,f"📎 Отправьте <b>файлом</b> прайс-лист поставщика <b>{supplier.name}</b>.\nPDF, Excel или другой документ будет прикреплён к поставщику.")
    return {"id":req.id,"status":"waiting"}


@router.delete("/price-lists/{price_list_id}")
def delete_price_list(price_list_id:int,user=Depends(get_current_user),db:Session=Depends(get_db)):
    obj=db.get(PriceList,price_list_id)
    if not obj:raise HTTPException(404,"Прайс-лист не найден")
    db.delete(obj);db.commit();return {"ok":True}


@router.get("/price-lists/{price_list_id}/content")
def price_list_content(price_list_id:int,user=Depends(get_current_user),db:Session=Depends(get_db)):
    obj=db.get(PriceList,price_list_id)
    if not obj:raise HTTPException(404,"Прайс-лист не найден")
    path=get_file_path(obj.telegram_file_id)
    url=file_download_url(path)
    client=httpx.Client(timeout=60)
    response=client.stream("GET",url);response.__enter__()
    if response.status_code!=200:
        response.__exit__(None,None,None);client.close();raise HTTPException(502,"Не удалось получить файл из Telegram")
    ctype=response.headers.get("content-type",obj.mime_type or "application/octet-stream")
    filename=obj.file_name or "price-list"
    def gen():
        try:
            for chunk in response.iter_bytes():yield chunk
        finally:
            response.__exit__(None,None,None);client.close()
    return StreamingResponse(gen(),media_type=ctype,headers={"Content-Disposition":f"attachment; filename*=UTF-8\'\'{quote(filename)}"})
