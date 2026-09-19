from __future__ import annotations

import html
import re
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path

import bleach
import httpx
from docx import Document
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.security import ALL_ROLES, get_current_user
from app.db.database import get_db
from app.db.models import (
    Employee,
    KnowledgeAcknowledgement,
    KnowledgeArticle,
    KnowledgeArticleLink,
    KnowledgeMedia,
    KnowledgeMediaUploadRequest,
    KnowledgeSection,
    Store,
    User,
)
from app.services.telegram import file_download_url, get_file_path
from app.services.notifications import notify_knowledge_published

router = APIRouter(prefix="/api/knowledge", tags=["knowledge"])

EDITOR_ROLES = {"admin", "leader", "operations_director"}
ARTICLE_STATUSES = {"draft", "published", "archived"}
MEDIA_KINDS = {"photo", "video"}
MAX_IMPORT_BYTES = 20 * 1024 * 1024
ALLOWED_IMPORT_EXTENSIONS = {".docx", ".pdf", ".txt", ".xlsx"}

ALLOWED_TAGS = [
    "p", "br", "h2", "h3", "h4", "strong", "b", "em", "i", "u", "s",
    "ul", "ol", "li", "blockquote", "hr", "a", "table", "thead", "tbody",
    "tr", "th", "td", "span", "div", "figure", "figcaption",
]
ALLOWED_ATTRS = {
    "a": ["href", "target", "rel"],
    "span": ["class", "data-media-id"],
    "div": ["class"],
    "figure": ["class", "data-media-id"],
    "td": ["colspan", "rowspan"],
    "th": ["colspan", "rowspan"],
}


def _require_editor(user: User):
    if user.role not in EDITOR_ROLES:
        raise HTTPException(403, "Редактировать базу знаний могут администратор, руководитель и операционный директор")


def _sanitize(raw: str | None) -> str:
    raw = re.sub(r"<\s*(script|style)[^>]*>.*?<\s*/\s*\1\s*>", "", raw or "", flags=re.I | re.S)
    return bleach.clean(
        raw,
        tags=ALLOWED_TAGS,
        attributes=ALLOWED_ATTRS,
        protocols={"http", "https", "mailto"},
        strip=True,
    )


def _plain(raw: str | None) -> str:
    if not raw:
        return ""
    # Keep spacing between blocks before stripping tags.
    text = re.sub(r"</(?:p|div|h[1-6]|li|tr|blockquote)>", " ", raw, flags=re.I)
    return re.sub(r"\s+", " ", bleach.clean(text, tags=[], strip=True)).strip()


def _read_minutes(raw: str | None) -> int:
    words = len(_plain(raw).split())
    return max(1, round(words / 180)) if words else 1


def _roles(section: KnowledgeSection) -> set[str]:
    return set(section.role_access or [])


def _section_chain(db: Session, section: KnowledgeSection) -> list[KnowledgeSection]:
    chain = [section]
    seen = {section.id}
    cur = section
    while cur.parent_id:
        cur = db.get(KnowledgeSection, cur.parent_id)
        if not cur or cur.id in seen:
            break
        chain.append(cur)
        seen.add(cur.id)
    return list(reversed(chain))


def _can_access_section(db: Session, user: User, section: KnowledgeSection) -> bool:
    if user.role in EDITOR_ROLES:
        return True
    for node in _section_chain(db, section):
        if not node.active:
            return False
        allowed = _roles(node)
        if allowed and user.role not in allowed:
            return False
    return True


def _ensure_section_access(db: Session, user: User, section_id: int) -> KnowledgeSection:
    section = db.get(KnowledgeSection, section_id)
    if not section:
        raise HTTPException(404, "Раздел не найден")
    if not _can_access_section(db, user, section):
        raise HTTPException(403, "Нет доступа к этому разделу")
    return section


def _section_path(db: Session, section: KnowledgeSection) -> list[dict]:
    return [{"id": x.id, "name": x.name} for x in _section_chain(db, section)]


def _media_dict(m: KnowledgeMedia) -> dict:
    return {
        "id": m.id,
        "article_id": m.article_id,
        "kind": m.kind,
        "file_name": m.file_name,
        "mime_type": m.mime_type,
        "file_size": m.file_size,
        "caption": m.caption,
        "created_at": m.created_at,
        "content_url": f"/api/knowledge/media/{m.id}/content",
    }


def _article_dict(db: Session, article: KnowledgeArticle, include_content: bool = True, user: User | None = None) -> dict:
    section = db.get(KnowledgeSection, article.section_id)
    author = db.get(User, article.author_id) if article.author_id else None
    store = db.get(Store, article.store_id) if article.store_id else None
    media = list(db.scalars(select(KnowledgeMedia).where(KnowledgeMedia.article_id == article.id).order_by(KnowledgeMedia.created_at)).all())
    acked = False
    if user:
        acked = db.scalar(select(KnowledgeAcknowledgement).where(
            KnowledgeAcknowledgement.article_id == article.id,
            KnowledgeAcknowledgement.user_id == user.id,
            KnowledgeAcknowledgement.revision == article.revision,
        )) is not None
    links = list(db.scalars(select(KnowledgeArticleLink).where(KnowledgeArticleLink.article_id == article.id).order_by(KnowledgeArticleLink.id)).all())
    payload = {
        "id": article.id,
        "section_id": article.section_id,
        "section_path": _section_path(db, section) if section else [],
        "title": article.title,
        "summary": article.summary,
        "status": article.status,
        "tags": article.tags_json or [],
        "store_id": article.store_id,
        "store_name": store.name if store else None,
        "position_tags": article.position_tags_json or [],
        "required_ack": article.required_ack,
        "revision": article.revision,
        "author_id": article.author_id,
        "author_name": author.full_name if author else None,
        "published_at": article.published_at,
        "created_at": article.created_at,
        "updated_at": article.updated_at,
        "read_minutes": _read_minutes(article.content_html),
        "acknowledged": acked,
        "media": [_media_dict(m) for m in media],
        "links": [{"id": x.id, "module_key": x.module_key, "entity_id": x.entity_id, "label": x.label} for x in links],
    }
    if include_content:
        payload["content_html"] = article.content_html or ""
    else:
        payload["snippet"] = (_plain(article.content_html)[:260] + "…") if len(_plain(article.content_html)) > 260 else _plain(article.content_html)
    return payload


def _article_visible(db: Session, user: User, article: KnowledgeArticle) -> bool:
    section = db.get(KnowledgeSection, article.section_id)
    if not section or not _can_access_section(db, user, section):
        return False
    if user.role in EDITOR_ROLES:
        return True
    return article.status == "published"


class SectionIn(BaseModel):
    parent_id: int | None = None
    name: str = Field(min_length=1, max_length=180)
    description: str | None = None
    icon: str | None = None
    role_access: list[str] = Field(default_factory=list)
    sort_order: int = 0
    active: bool = True


class ArticleIn(BaseModel):
    section_id: int
    title: str = Field(min_length=1, max_length=300)
    summary: str | None = None
    content_html: str = ""
    status: str = "draft"
    tags: list[str] = Field(default_factory=list)
    store_id: int | None = None
    position_tags: list[str] = Field(default_factory=list)
    required_ack: bool = False
    links: list[dict] = Field(default_factory=list)


class MediaRequestIn(BaseModel):
    kind: str


class MediaPatch(BaseModel):
    caption: str | None = None


@router.get("/sections")
def list_sections(manage: bool = False, user=Depends(get_current_user), db: Session = Depends(get_db)):
    if manage:
        _require_editor(user)
    rows = list(db.scalars(select(KnowledgeSection).order_by(KnowledgeSection.sort_order, KnowledgeSection.name)).all())
    result = []
    for s in rows:
        if not manage and (not s.active or not _can_access_section(db, user, s)):
            continue
        count_q = select(KnowledgeArticle).where(KnowledgeArticle.section_id == s.id)
        if not manage:
            count_q = count_q.where(KnowledgeArticle.status == "published")
        article_count = len(list(db.scalars(count_q).all()))
        result.append({
            "id": s.id,
            "parent_id": s.parent_id,
            "name": s.name,
            "description": s.description,
            "icon": s.icon,
            "role_access": s.role_access or [],
            "sort_order": s.sort_order,
            "active": s.active,
            "article_count": article_count,
        })
    return result


@router.post("/sections")
def create_section(payload: SectionIn, user=Depends(get_current_user), db: Session = Depends(get_db)):
    _require_editor(user)
    if payload.parent_id:
        parent = db.get(KnowledgeSection, payload.parent_id)
        if not parent:
            raise HTTPException(400, "Родительский раздел не найден")
        if parent.parent_id:
            raise HTTPException(400, "Поддерживается структура: раздел → подраздел → статья")
    invalid = set(payload.role_access) - ALL_ROLES
    if invalid:
        raise HTTPException(400, f"Неизвестные роли: {', '.join(sorted(invalid))}")
    exists = db.scalar(select(KnowledgeSection).where(KnowledgeSection.parent_id == payload.parent_id, KnowledgeSection.name == payload.name.strip()))
    if exists:
        raise HTTPException(400, "Раздел с таким названием уже существует")
    data = payload.model_dump()
    data["name"] = payload.name.strip()
    s = KnowledgeSection(**data, created_by=user.id, updated_by=user.id)
    db.add(s); db.commit(); db.refresh(s)
    return {"id": s.id}


@router.patch("/sections/{section_id}")
def update_section(section_id: int, payload: dict, user=Depends(get_current_user), db: Session = Depends(get_db)):
    _require_editor(user)
    s = db.get(KnowledgeSection, section_id)
    if not s:
        raise HTTPException(404, "Раздел не найден")
    allowed = {"name", "description", "icon", "role_access", "sort_order", "active"}
    for k, v in payload.items():
        if k not in allowed:
            continue
        if k == "role_access":
            invalid = set(v or []) - ALL_ROLES
            if invalid:
                raise HTTPException(400, "Есть неизвестные роли")
        setattr(s, k, v)
    s.updated_by = user.id
    db.commit()
    return {"ok": True}


@router.delete("/sections/{section_id}")
def delete_section(section_id: int, user=Depends(get_current_user), db: Session = Depends(get_db)):
    _require_editor(user)
    s = db.get(KnowledgeSection, section_id)
    if not s:
        raise HTTPException(404, "Раздел не найден")
    child = db.scalar(select(KnowledgeSection).where(KnowledgeSection.parent_id == s.id))
    article = db.scalar(select(KnowledgeArticle).where(KnowledgeArticle.section_id == s.id))
    if child or article:
        s.active = False
        s.updated_by = user.id
        db.commit()
        return {"ok": True, "deactivated": True}
    db.delete(s); db.commit()
    return {"ok": True, "deleted": True}


@router.get("/articles")
def list_articles(section_id: int | None = None, status: str | None = None, user=Depends(get_current_user), db: Session = Depends(get_db)):
    q = select(KnowledgeArticle).order_by(KnowledgeArticle.updated_at.desc())
    if section_id:
        q = q.where(KnowledgeArticle.section_id == section_id)
    if status and user.role in EDITOR_ROLES:
        q = q.where(KnowledgeArticle.status == status)
    elif user.role not in EDITOR_ROLES:
        q = q.where(KnowledgeArticle.status == "published")
    rows = []
    for a in db.scalars(q).all():
        if _article_visible(db, user, a):
            rows.append(_article_dict(db, a, include_content=False, user=user))
    return rows


@router.get("/articles/{article_id}")
def get_article(article_id: int, user=Depends(get_current_user), db: Session = Depends(get_db)):
    a = db.get(KnowledgeArticle, article_id)
    if not a or not _article_visible(db, user, a):
        raise HTTPException(404, "Статья не найдена или недоступна")
    return _article_dict(db, a, include_content=True, user=user)


def _replace_links(db: Session, article_id: int, links: list[dict]):
    for old in list(db.scalars(select(KnowledgeArticleLink).where(KnowledgeArticleLink.article_id == article_id)).all()):
        db.delete(old)
    for x in links or []:
        key = str(x.get("module_key") or "").strip()
        if not key:
            continue
        db.add(KnowledgeArticleLink(article_id=article_id, module_key=key[:64], entity_id=x.get("entity_id"), label=(x.get("label") or "")[:180] or None))


@router.post("/articles")
def create_article(payload: ArticleIn, user=Depends(get_current_user), db: Session = Depends(get_db)):
    _require_editor(user)
    section = db.get(KnowledgeSection, payload.section_id)
    if not section:
        raise HTTPException(400, "Раздел не найден")
    if payload.status not in ARTICLE_STATUSES:
        raise HTTPException(400, "Некорректный статус")
    a = KnowledgeArticle(
        section_id=payload.section_id,
        title=payload.title.strip(),
        summary=payload.summary,
        content_html=_sanitize(payload.content_html),
        status=payload.status,
        tags_json=[x.strip() for x in payload.tags if x.strip()],
        store_id=payload.store_id,
        position_tags_json=payload.position_tags,
        required_ack=payload.required_ack,
        revision=1,
        author_id=user.id,
        updated_by=user.id,
        published_at=datetime.utcnow() if payload.status == "published" else None,
    )
    db.add(a); db.flush()
    _replace_links(db, a.id, payload.links)
    if a.status == "published":
        try:
            notify_knowledge_published(db, a.id)
        except Exception:
            pass
    db.commit(); db.refresh(a)
    return _article_dict(db, a, user=user)


@router.patch("/articles/{article_id}")
def update_article(article_id: int, payload: dict, user=Depends(get_current_user), db: Session = Depends(get_db)):
    _require_editor(user)
    a = db.get(KnowledgeArticle, article_id)
    if not a:
        raise HTTPException(404, "Статья не найдена")
    old_status = a.status
    content_keys = {"section_id", "title", "summary", "content_html", "tags", "store_id", "position_tags", "required_ack"}
    content_changed = False
    if "section_id" in payload and not db.get(KnowledgeSection, int(payload["section_id"])):
        raise HTTPException(400, "Раздел не найден")
    if "status" in payload and payload["status"] not in ARTICLE_STATUSES:
        raise HTTPException(400, "Некорректный статус")
    mapping = {"tags": "tags_json", "position_tags": "position_tags_json"}
    for key in ["section_id", "title", "summary", "content_html", "status", "tags", "store_id", "position_tags", "required_ack"]:
        if key not in payload:
            continue
        value = payload[key]
        if key == "content_html":
            value = _sanitize(value)
        elif key == "title":
            value = str(value).strip()
        elif key in {"tags", "position_tags"}:
            value = [str(x).strip() for x in (value or []) if str(x).strip()]
        attr = mapping.get(key, key)
        if getattr(a, attr) != value:
            if key in content_keys:
                content_changed = True
            setattr(a, attr, value)
    if payload.get("status") == "published" and not a.published_at:
        a.published_at = datetime.utcnow()
    if content_changed:
        a.revision = int(a.revision or 1) + 1
    if "links" in payload:
        _replace_links(db, a.id, payload.get("links") or [])
    a.updated_by = user.id
    if old_status != "published" and a.status == "published":
        try:
            notify_knowledge_published(db, a.id)
        except Exception:
            pass
    db.commit(); db.refresh(a)
    return _article_dict(db, a, user=user)


@router.delete("/articles/{article_id}")
def archive_article(article_id: int, user=Depends(get_current_user), db: Session = Depends(get_db)):
    _require_editor(user)
    a = db.get(KnowledgeArticle, article_id)
    if not a:
        raise HTTPException(404, "Статья не найдена")
    a.status = "archived"
    a.updated_by = user.id
    db.commit()
    return {"ok": True}


@router.get("/search")
def search_knowledge(q: str, user=Depends(get_current_user), db: Session = Depends(get_db)):
    query = q.strip().lower()
    if not query:
        return []
    # Deliberately bounded and permission-filtered; later this endpoint can be backed by PostgreSQL FTS.
    candidates = list(db.scalars(select(KnowledgeArticle).where(KnowledgeArticle.status == "published").order_by(KnowledgeArticle.updated_at.desc()).limit(500)).all())
    scored = []
    for a in candidates:
        if not _article_visible(db, user, a):
            continue
        section = db.get(KnowledgeSection, a.section_id)
        section_text = " ".join(x["name"] for x in _section_path(db, section)) if section else ""
        tags = " ".join(a.tags_json or [])
        body = _plain(a.content_html)
        hay = f"{a.title} {a.summary or ''} {section_text} {tags} {body}".lower()
        if query not in hay:
            continue
        score = 0
        if query in a.title.lower(): score += 5
        if query in tags.lower(): score += 3
        if query in section_text.lower(): score += 2
        if query in body.lower(): score += 1
        item = _article_dict(db, a, include_content=False, user=user)
        item["search_score"] = score
        scored.append(item)
    scored.sort(key=lambda x: (x["search_score"], x["updated_at"] or datetime.min), reverse=True)
    return scored[:50]


@router.get("/required")
def required_for_me(user=Depends(get_current_user), db: Session = Depends(get_db)):
    rows = list(db.scalars(select(KnowledgeArticle).where(KnowledgeArticle.status == "published", KnowledgeArticle.required_ack.is_(True)).order_by(KnowledgeArticle.published_at.desc())).all())
    result = []
    for a in rows:
        if not _article_visible(db, user, a):
            continue
        ack = db.scalar(select(KnowledgeAcknowledgement).where(
            KnowledgeAcknowledgement.article_id == a.id,
            KnowledgeAcknowledgement.user_id == user.id,
            KnowledgeAcknowledgement.revision == a.revision,
        ))
        item = _article_dict(db, a, include_content=False, user=user)
        item["pending"] = ack is None
        result.append(item)
    return result


@router.post("/articles/{article_id}/acknowledge")
def acknowledge(article_id: int, user=Depends(get_current_user), db: Session = Depends(get_db)):
    a = db.get(KnowledgeArticle, article_id)
    if not a or a.status != "published" or not _article_visible(db, user, a):
        raise HTTPException(404, "Статья недоступна")
    if not a.required_ack:
        raise HTTPException(400, "Для этой статьи подтверждение ознакомления не требуется")
    existing = db.scalar(select(KnowledgeAcknowledgement).where(
        KnowledgeAcknowledgement.article_id == a.id,
        KnowledgeAcknowledgement.user_id == user.id,
        KnowledgeAcknowledgement.revision == a.revision,
    ))
    if existing:
        return {"ok": True, "acknowledged_at": existing.acknowledged_at}
    employee = db.scalar(select(Employee).where(Employee.user_id == user.id))
    row = KnowledgeAcknowledgement(article_id=a.id, user_id=user.id, employee_id=employee.id if employee else None, revision=a.revision)
    db.add(row); db.flush()
    # Linked tests may be assigned automatically only after the employee explicitly confirms acknowledgement.
    try:
        from app.services.testing import ensure_article_ack_assignments
        assigned_test_ids = ensure_article_ack_assignments(db, a.id, user.id)
    except Exception:
        assigned_test_ids = []
    db.commit(); db.refresh(row)
    return {"ok": True, "acknowledged_at": row.acknowledged_at, "test_assignments": assigned_test_ids}


@router.get("/acknowledgements/control")
def acknowledgement_control(article_id: int, user=Depends(get_current_user), db: Session = Depends(get_db)):
    _require_editor(user)
    a = db.get(KnowledgeArticle, article_id)
    if not a:
        raise HTTPException(404, "Статья не найдена")
    section = db.get(KnowledgeSection, a.section_id)
    if not section:
        return []
    effective_roles = set(ALL_ROLES)
    for node in _section_chain(db, section):
        if node.role_access:
            effective_roles &= set(node.role_access)
    employees = list(db.scalars(select(Employee).where(Employee.active.is_(True)).order_by(Employee.full_name)).all())
    result = []
    for e in employees:
        if effective_roles and e.position not in effective_roles:
            continue
        ack = None
        if e.user_id:
            ack = db.scalar(select(KnowledgeAcknowledgement).where(
                KnowledgeAcknowledgement.article_id == a.id,
                KnowledgeAcknowledgement.user_id == e.user_id,
                KnowledgeAcknowledgement.revision == a.revision,
            ))
        result.append({
            "employee_id": e.id,
            "full_name": e.full_name,
            "position": e.position,
            "has_account": bool(e.user_id),
            "acknowledged": bool(ack),
            "acknowledged_at": ack.acknowledged_at if ack else None,
        })
    return result


@router.post("/articles/{article_id}/media-request")
def media_request(article_id: int, payload: MediaRequestIn, user=Depends(get_current_user), db: Session = Depends(get_db)):
    _require_editor(user)
    if payload.kind not in MEDIA_KINDS:
        raise HTTPException(400, "Поддерживаются фото и видео")
    a = db.get(KnowledgeArticle, article_id)
    if not a:
        raise HTTPException(404, "Сначала сохраните статью")
    req = KnowledgeMediaUploadRequest(article_id=a.id, user_id=user.id, kind=payload.kind, expires_at=datetime.utcnow() + timedelta(minutes=30))
    db.add(req); db.commit(); db.refresh(req)
    return {"ok": True, "request_id": req.id, "message": "Отправьте медиа в чат с ботом"}


@router.patch("/media/{media_id}")
def update_media(media_id: int, payload: MediaPatch, user=Depends(get_current_user), db: Session = Depends(get_db)):
    _require_editor(user)
    m = db.get(KnowledgeMedia, media_id)
    if not m:
        raise HTTPException(404, "Медиа не найдено")
    m.caption = payload.caption
    db.commit()
    return {"ok": True}


@router.delete("/media/{media_id}")
def delete_media(media_id: int, user=Depends(get_current_user), db: Session = Depends(get_db)):
    _require_editor(user)
    m = db.get(KnowledgeMedia, media_id)
    if not m:
        raise HTTPException(404, "Медиа не найдено")
    db.delete(m); db.commit()
    return {"ok": True}


@router.get("/media/{media_id}/content")
def media_content(media_id: int, user=Depends(get_current_user), db: Session = Depends(get_db)):
    m = db.get(KnowledgeMedia, media_id)
    if not m:
        raise HTTPException(404, "Медиа не найдено")
    article = db.get(KnowledgeArticle, m.article_id)
    if not article or not _article_visible(db, user, article):
        raise HTTPException(403, "Нет доступа к медиа")
    try:
        path = get_file_path(m.telegram_file_id)
        url = file_download_url(path)
        with httpx.Client(timeout=60) as client:
            r = client.get(url)
            r.raise_for_status()
            return Response(content=r.content, media_type=m.mime_type or ("video/mp4" if m.kind == "video" else "image/jpeg"), headers={"Cache-Control": "private, max-age=300"})
    except Exception as e:
        raise HTTPException(502, f"Не удалось получить файл из Telegram: {e}")


def _docx_to_html(data: bytes) -> str:
    doc = Document(BytesIO(data))
    chunks = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style else ""
        tag = "h2" if "heading 1" in style or "заголовок 1" in style else "h3" if "heading 2" in style or "заголовок 2" in style else "p"
        chunks.append(f"<{tag}>{html.escape(text)}</{tag}>")
    for table in doc.tables:
        chunks.append("<table><tbody>")
        for row in table.rows:
            chunks.append("<tr>" + "".join(f"<td>{html.escape(cell.text)}</td>" for cell in row.cells) + "</tr>")
        chunks.append("</tbody></table>")
    return "".join(chunks)


def _pdf_to_html(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    chunks = []
    for idx, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            chunks.append(f"<h3>Страница {idx}</h3>" + "".join(f"<p>{html.escape(line)}</p>" for line in text.splitlines() if line.strip()))
    if not chunks:
        raise HTTPException(400, "В PDF не найден извлекаемый текст. Сканированные PDF пока нужно преобразовать в текстовый PDF или DOCX")
    return "".join(chunks)


def _txt_to_html(data: bytes) -> str:
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            pass
    if text is None:
        raise HTTPException(400, "Не удалось определить кодировку TXT")
    return "".join(f"<p>{html.escape(line)}</p>" for line in text.splitlines() if line.strip())


def _xlsx_to_html(data: bytes) -> str:
    wb = load_workbook(BytesIO(data), data_only=True, read_only=True)
    chunks = []
    for ws in wb.worksheets:
        chunks.append(f"<h2>{html.escape(ws.title)}</h2><table><tbody>")
        for row in ws.iter_rows(values_only=True):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            chunks.append("<tr>" + "".join(f"<td>{html.escape('' if v is None else str(v))}</td>" for v in row) + "</tr>")
        chunks.append("</tbody></table>")
    return "".join(chunks)


@router.post("/import")
async def import_article(
    section_id: int = Form(...),
    title: str | None = Form(None),
    file: UploadFile = File(...),
    user=Depends(get_current_user),
    db: Session = Depends(get_db),
):
    _require_editor(user)
    section = db.get(KnowledgeSection, section_id)
    if not section:
        raise HTTPException(400, "Раздел не найден")
    name = file.filename or "document"
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED_IMPORT_EXTENSIONS:
        raise HTTPException(400, "Поддерживаются DOCX, PDF, TXT и XLSX")
    data = await file.read(MAX_IMPORT_BYTES + 1)
    if len(data) > MAX_IMPORT_BYTES:
        raise HTTPException(400, "Файл больше 20 МБ")
    if ext == ".docx": content = _docx_to_html(data)
    elif ext == ".pdf": content = _pdf_to_html(data)
    elif ext == ".txt": content = _txt_to_html(data)
    else: content = _xlsx_to_html(data)
    article = KnowledgeArticle(
        section_id=section_id,
        title=(title or Path(name).stem).strip()[:300],
        content_html=_sanitize(content),
        status="draft",
        tags_json=[],
        position_tags_json=[],
        required_ack=False,
        revision=1,
        author_id=user.id,
        updated_by=user.id,
    )
    db.add(article); db.commit(); db.refresh(article)
    return _article_dict(db, article, user=user)


@router.get("/related")
def related_articles(module_key: str, entity_id: int | None = None, user=Depends(get_current_user), db: Session = Depends(get_db)):
    q = select(KnowledgeArticleLink).where(KnowledgeArticleLink.module_key == module_key)
    if entity_id is not None:
        q = q.where((KnowledgeArticleLink.entity_id == entity_id) | (KnowledgeArticleLink.entity_id.is_(None)))
    else:
        q = q.where(KnowledgeArticleLink.entity_id.is_(None))
    result = []
    for link in db.scalars(q).all():
        a = db.get(KnowledgeArticle, link.article_id)
        if a and a.status == "published" and _article_visible(db, user, a):
            item = _article_dict(db, a, include_content=False, user=user)
            item["link_label"] = link.label
            result.append(item)
    return result


@router.get("/context-search")
def ai_context_search(q: str, user=Depends(get_current_user), db: Session = Depends(get_db)):
    """Permission-aware knowledge context for the future AI agent."""
    results = search_knowledge(q=q, user=user, db=db)[:10]
    out = []
    for item in results:
        a = db.get(KnowledgeArticle, item["id"])
        out.append({
            "article_id": a.id,
            "title": a.title,
            "section_path": item["section_path"],
            "tags": a.tags_json or [],
            "text": _plain(a.content_html)[:7000],
            "updated_at": a.updated_at,
        })
    return {"query": q, "articles": out}
